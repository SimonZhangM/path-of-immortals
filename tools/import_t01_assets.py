"""Read current T01-08 category rows; preserve existing game identities/art.

Run from project root with Python + openpyxl. Never rewrites the source workbook.
Effect rows are cross-checked references, not a second source of damage packets.
"""
import hashlib
import json
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'output/spreadsheet/t01-08-物品资产主表-260928.xlsx'
TARGET = ROOT / 'data/maps/inventory_items.json'
ELEMENTS = dict(zip(['NONE', '金', '木', '水', '火', '土', '风', '雷'],
                    ['none', 'metal', 'wood', 'water', 'fire', 'earth', 'wind', 'thunder']))
BUFFS = ['反锋', '生机', '润脉', '附炎', '坚韧', '轻灵', '雷蕴']


def export():
    book = openpyxl.load_workbook(SOURCE, data_only=True)
    formulas = openpyxl.load_workbook(SOURCE, data_only=False)
    old = {r['id']: r for r in json.loads(TARGET.read_text(encoding='utf-8-sig'))}
    records, mapping = [], []
    effects_by_asset = {}
    for row in list(book['17_效果与触发'].values)[5:]:
        effects_by_asset.setdefault(row[1], []).append(row[0])
    for sheet, category in [('02_武器', 'weapon'), ('03_防具', 'armor'),
                            ('04_法器', 'artifact'), ('05_丹药', 'pill'), ('07_暗器', 'throwable')]:
        ws = book[sheet]
        headers = [c.value for c in ws[5]]
        for row in ws.iter_rows(min_row=6):
            d = dict(zip(headers, [c.value for c in row]))
            asset = d['资产ID']
            if not isinstance(asset, str) or not asset.startswith('t01.asset.'):
                continue
            for cell in row:
                if formulas[sheet][cell.coordinate].data_type == 'f' and cell.value is None:
                    raise ValueError(f'Missing formula cache: {sheet}!{cell.coordinate}')
            existing = d.get('现有游戏ID')
            game_id = existing if isinstance(existing, str) and existing.startswith('base.') else 'base.map_item.t01_' + asset.rsplit('.', 1)[1]
            r = dict(old.get(game_id, {}))
            r.update(id=game_id, asset_id=asset, name=d['名称'], category=category, quality=d['品级'],
                     quantity=r.get('quantity', 1), acquired_at=r.get('acquired_at', int(asset.rsplit('.', 1)[1]) + 100),
                     enhancement_level=r.get('enhancement_level', 0), icon=r.get('icon', ''),
                     card_frame=r.get('card_frame', 'res://assets/level-1.webp' if d['品级'] == '下品' else 'res://assets/level-2.webp'),
                     footprint_columns=d['宽_格'], footprint_rows=d['高_格'],
                     element='base.element.' + ELEMENTS[d['元素']], rule_version=1,
                     effects=[], base_stamina_cost=0, spirit_cost=0, cooldown=d.get('CD_秒') or d.get('设计CD_秒') or 0,
                     source_location=f'{SOURCE.name}::{sheet}!A{row[0].row}', uses_per_unit=1 if category in ['pill', 'throwable'] else 0)
            def effect(kind, **kw):
                r['effects'].append(dict(trigger='on_activate', effect=kind, **kw))
            if category == 'weapon':
                r.update(subcategory=d['器型'], damage_type=d['Attack Type'], base_damage=d['设计伤害'],
                         base_stamina_cost=d['设计体力费'], spirit_cost=d.get('灵力费') or 0, hit_chance=0.9)
                effect('damage', value=r['base_damage'], damage_type=r['damage_type'])
                if d.get('状态'):
                    effect('apply_status', status=d['状态'], value=d['施加层数'],
                           target='self' if d['状态'] in BUFFS else 'enemy',
                           gate='hp_damage' if '实际伤到气血' in d['状态时机与目标'] else 'hit')
            elif category == 'armor':
                r.update(subcategory=d['部位'], armor_type=d['甲型'], armor_capacity=d['护甲上限贡献'], armor_gain=d['单次回甲'] or 0)
                if r['armor_gain']:
                    effect('restore_armor', value=r['armor_gain'])
                if d.get('状态'):
                    effect('apply_status', status=d['状态'], value=d['层数'], target='self', gate='always')
                    if d['状态触发'] == '受击':
                        r['effects'][-1].update(trigger='on_attacked', cooldown=9)
            elif category == 'artifact':
                r.update(subcategory='法器', spirit_cost=d.get('灵力费') or 0, base_stamina_cost=d.get('体力费') or 0,
                         armor_capacity=d.get('护甲上限贡献') or 0)
                for cn, resource in [('气血恢复', 'hp'), ('体力恢复', 'stamina'), ('灵力恢复', 'spirit')]:
                    if d[cn]:
                        effect('restore_capped', resource=resource, value=d[cn], cap_numerator=d['恢复上限分子'], cap_denominator=d['恢复上限分母'])
                if d.get('本体伤害'):
                    r.update(damage_type=d['攻击类型'], base_damage=d['本体伤害'], hit_chance=0.95)
                    effect('damage', value=d['本体伤害'], damage_type=d['攻击类型'])
                if d.get('单次回甲'):
                    effect('restore_armor', value=d['单次回甲'])
                if d.get('普通灵盾生成'):
                    effect('barrier', value=d['普通灵盾生成'])
                    r['barrier_full_stop'] = True
                if d.get('状态'):
                    effect('apply_status', status=d['状态'], value=d['层数'], target='self' if d['状态'] in BUFFS else 'enemy',
                           gate='hit' if '命中' in d['施加时机与目标'] else 'always')
            elif category == 'pill':
                r.update(subcategory='丹药', first_ready=True)
                for cn, resource in [('每次气血恢复', 'hp'), ('每次体力恢复', 'stamina'), ('每次灵力恢复', 'spirit')]:
                    if d[cn]:
                        if d['跳间隔_秒']:
                            effect('restore_ticks', resource=resource, value=d[cn], ticks=d['次数'], interval=d['跳间隔_秒'])
                        else:
                            effect('restore_instant', resource=resource, value=d[cn])
                if d['毒蚀免疫_秒']:
                    effect('cleanse_toxin', duration=d['毒蚀免疫_秒'])
                if d.get('状态'):
                    effect('apply_status', status=d['状态'], value=d['层数'], target='self', gate='always')
                if d.get('减伤元素'):
                    effect('resistance', element='base.element.' + ELEMENTS[d['减伤元素']], value=d['减伤比例'], duration=d['持续_秒'])
            else:
                r.update(subcategory='暗器', damage_type=d['攻击类型'], base_damage=d['基础伤害'], hit_chance=0.9)
                effect('damage', value=d['基础伤害'], damage_type=d['攻击类型'])
                if d.get('状态'):
                    effect('apply_status', status=d['状态'], value=d['层数'], target='enemy', gate='hit')
            dependencies = sorted(set(e.get('status', e['effect']) for e in r['effects']))
            refs = effects_by_asset.get(asset, [])
            if d.get('效果引用'):
                assert set(d['效果引用'].split('；')) == set(refs), (asset, refs)
            records.append(r)
            mapping.append(dict(asset_id=asset, name=r['name'], quality=r['quality'], game_id=game_id,
                                source=r['source_location'], effects=refs, dependencies=dependencies,
                                previous='需迁移' if existing and existing.startswith('base.') else '需新增',
                                implementation='定义已转换；行为与入口待验证', missing='命中模板' if r.get('hit_chance', 1) is None else ''))
    assert len(records) == 115 and len({r['id'] for r in records}) == 115
    assert sum(r['quality'] == '下品' for r in records) == 27
    # Explicit member mappings; no recipe/cost inference.
    by_asset = {r['asset_id']: r for r in records}
    for row in list(book['23_联动升品映射'].values)[5:]:
        for asset in [row[3], row[5]]:
            if asset in by_asset:
                by_asset[asset]['lineage'] = row[1]
    unrelated = [r for r in old.values() if r['id'] not in {r['id'] for r in records}]
    TARGET.write_text(json.dumps(records + unrelated, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    out = ROOT / 'docs/sources/t01_asset_implementation.json'
    out.write_text(json.dumps(dict(source=SOURCE.name, sha256=hashlib.sha256(SOURCE.read_bytes()).hexdigest(), assets=mapping), ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(f'Exported {len(records)} T01 assets; preserved {len(unrelated)} unrelated records.')


if __name__ == '__main__':
    export()
